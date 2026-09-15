"""Presentation parser for PPTX and PDF courseware."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
import pymupdf


@dataclass
class SlideInfo:
    """Represents a single parsed slide from a presentation."""
    index: int  # 1-based index
    title: str = ""
    bullets: list[str] = field(default_factory=list)
    raw_text: str = ""
    notes: str = ""
    narration: str = ""
    layout: str = "pip"  # pip, split, full_avatar, full_slide
    image_path: Path | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "title": self.title,
            "bullets": self.bullets,
            "raw_text": self.raw_text,
            "notes": self.notes,
            "narration": self.narration,
            "layout": self.layout,
            "image_path": str(self.image_path) if self.image_path else None,
        }


@dataclass
class CourseDeck:
    """Represents an entire parsed presentation deck."""
    title: str
    source_file: Path
    slides: list[SlideInfo] = field(default_factory=list)

    @property
    def total_slides(self) -> int:
        return len(self.slides)

    def to_dict(self) -> dict[str, Any]:
        return {
            "title": self.title,
            "source_file": str(self.source_file),
            "total_slides": self.total_slides,
            "slides": [s.to_dict() for s in self.slides],
        }


def extract_layout_directive(text: str) -> tuple[str | None, str]:
    """Extract optional layout tag like [layout: pip] or [布局: 分屏] and return (layout, cleaned_text)."""
    match = re.search(r"\[(?:layout|布局)\s*[:=：]\s*([a-zA-Z_]+|画中画|分屏|全屏讲师|全屏课件)\]", text, re.IGNORECASE)
    if not match:
        return None, text
    raw_tag = match.group(1).lower().strip()
    tag_map = {
        "pip": "pip",
        "画中画": "pip",
        "split": "split",
        "分屏": "split",
        "full_avatar": "full_avatar",
        "全屏讲师": "full_avatar",
        "full_slide": "full_slide",
        "全屏课件": "full_slide",
    }
    layout = tag_map.get(raw_tag)
    cleaned = re.sub(r"\[(?:layout|布局)\s*[:=：]\s*([a-zA-Z_]+|画中画|分屏|全屏讲师|全屏课件)\]", "", text, flags=re.IGNORECASE).strip()
    return layout, cleaned


class PresentationParser:
    """Extracts slides, text, speaker notes, and layouts from PPTX or PDF."""

    @classmethod
    def parse(cls, file_path: str | Path) -> CourseDeck:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"presentation file not found: {path}")

        suffix = path.suffix.lower()
        if suffix in {".pptx", ".ppt"}:
            return cls._parse_pptx(path)
        elif suffix == ".pdf":
            return cls._parse_pdf(path)
        else:
            raise ValueError(f"unsupported presentation format: {suffix}, must be .pptx or .pdf")

    @classmethod
    def _parse_pptx(cls, path: Path) -> CourseDeck:
        prs = Presentation(str(path))
        deck_title = path.stem
        slides: list[SlideInfo] = []

        total = len(prs.slides)
        for idx, slide in enumerate(prs.slides, start=1):
            title = ""
            bullets: list[str] = []
            text_parts: list[str] = []

            # Extract title and body text
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()
                if not deck_title or deck_title == path.stem:
                    deck_title = title

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                # Skip the title shape if we already recorded it
                if shape == slide.shapes.title:
                    continue
                for p in shape.text_frame.paragraphs:
                    line = p.text.strip()
                    if line:
                        bullets.append(line)
                        text_parts.append(line)

            raw_text = "\n".join(text_parts)

            # Extract speaker notes
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            # Determine narration and layout
            directive_layout, cleaned_notes = extract_layout_directive(notes)
            narration = cleaned_notes if cleaned_notes else (raw_text if raw_text else title)

            # Default layout: Standard micro-course layout with slide & instructor
            layout = directive_layout if directive_layout else "pip"

            slides.append(
                SlideInfo(
                    index=idx,
                    title=title or f"第 {idx} 页",
                    bullets=bullets,
                    raw_text=raw_text,
                    notes=notes,
                    narration=narration,
                    layout=layout,
                )
            )

        return CourseDeck(title=deck_title or path.stem, source_file=path, slides=slides)

    @classmethod
    def _parse_pdf(cls, path: Path) -> CourseDeck:
        doc = pymupdf.open(str(path))
        deck_title = path.stem
        slides: list[SlideInfo] = []
        total = len(doc)

        for idx, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            title = lines[0] if lines else f"第 {idx} 页"
            bullets = lines[1:] if len(lines) > 1 else []
            raw_text = "\n".join(bullets)

            directive_layout, cleaned_text = extract_layout_directive(raw_text)
            narration = cleaned_text if cleaned_text else title
            layout = directive_layout if directive_layout else "pip"

            slides.append(
                SlideInfo(
                    index=idx,
                    title=title,
                    bullets=bullets,
                    raw_text=raw_text,
                    notes="",
                    narration=narration,
                    layout=layout,
                )
            )

        return CourseDeck(title=deck_title, source_file=path, slides=slides)
