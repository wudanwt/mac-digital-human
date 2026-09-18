from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from pptx import Presentation

from app.saas_main import app


def _register(client: TestClient, prefix: str) -> tuple[str, str, str]:
    email = f"{prefix}-{uuid4().hex[:8]}@example.com"
    password = "correct-horse-battery-staple"
    response = client.post(
        "/api/saas/auth/register",
        json={
            "email": email,
            "password": password,
            "display_name": prefix,
            "workspace_name": f"{prefix} workspace",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"], email, password


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "第一章：市场变化"
    slide.placeholders[1].text = "客户需求从项目交付转向持续经营"
    slide.notes_slide.notes_text_frame.text = "这一页重点解释客户为什么需要持续经营能力。"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_account_profile_and_password_change() -> None:
    with TestClient(app) as client:
        token, email, password = _register(client, "account-tools")
        headers = _headers(token)

        account = client.get("/api/saas/account", headers=headers)
        assert account.status_code == 200
        assert account.json()["email"] == email

        updated = client.patch(
            "/api/saas/account",
            headers=headers,
            json={"display_name": "新的显示名称"},
        )
        assert updated.status_code == 200
        assert updated.json()["display_name"] == "新的显示名称"

        new_password = "new-correct-horse-battery-staple"
        changed = client.post(
            "/api/saas/account/change-password",
            headers=headers,
            json={"current_password": password, "new_password": new_password},
        )
        assert changed.status_code == 204, changed.text

        old_login = client.post("/api/saas/auth/login", json={"email": email, "password": password})
        assert old_login.status_code == 401
        new_login = client.post("/api/saas/auth/login", json={"email": email, "password": new_password})
        assert new_login.status_code == 200


def test_ppt_outline_parser_and_thumbnail_for_course_studio() -> None:
    with TestClient(app) as client:
        token, _, _ = _register(client, "course-tools")
        headers = _headers(token)
        upload = client.post(
            "/api/saas/assets",
            headers=headers,
            data={"kind": "ppt"},
            files={
                "file": (
                    "training.pptx",
                    BytesIO(_pptx_bytes()),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        assert upload.status_code == 201, upload.text
        asset_id = upload.json()["id"]

        outline = client.get(f"/api/saas/course-tools/ppt/{asset_id}/outline", headers=headers)
        assert outline.status_code == 200, outline.text
        data = outline.json()
        assert data["total_slides"] == 1
        assert data["slides"][0]["title"] == "第一章：市场变化"
        assert "持续经营" in data["slides"][0]["narration"]
        assert data["slides"][0]["thumbnail_url"].endswith("/slides/1/thumbnail")

        thumbnail = client.get(data["slides"][0]["thumbnail_url"], headers=headers)
        assert thumbnail.status_code == 200, thumbnail.text
        assert thumbnail.headers["content-type"].startswith("image/png")
        assert len(thumbnail.content) > 100


def test_builtin_background_catalog_and_preview() -> None:
    with TestClient(app) as client:
        token, _, _ = _register(client, "background-tools")
        headers = _headers(token)
        response = client.get("/api/saas/course-tools/backgrounds", headers=headers)
        assert response.status_code == 200, response.text
        themes = response.json()
        assert len(themes) >= 8
        assert {item["id"] for item in themes} >= {
            "deep-space-grid",
            "aurora-cyan",
            "executive-blue",
            "energy-network",
        }
        preview = client.get(themes[0]["preview_url"], headers=headers)
        assert preview.status_code == 200
        assert preview.headers["content-type"].startswith("image/png")
        assert preview.content.startswith(b"\x89PNG")


def test_theme_and_course_studio_assets_are_served() -> None:
    with TestClient(app) as client:
        css = client.get("/saas-theme.css")
        studio_css = client.get("/course-studio.css")
        studio_upgrade_css = client.get("/course-studio-upgrade.css")
        product_js = client.get("/product-ui.js")
        polish_js = client.get("/polish-ui.js")
        studio_js = client.get("/course-studio.js")
        studio_upgrade_js = client.get("/course-studio-upgrade.js")
        page = client.get("/")
        assert css.status_code == 200
        assert "tech-hero" in css.text
        assert studio_css.status_code == 200
        assert "course-studio-shell" in studio_css.text
        assert studio_upgrade_css.status_code == 200
        assert "premium-narration" in studio_upgrade_css.text
        assert "studio-bg-grid" in studio_upgrade_css.text
        assert product_js.status_code == 200
        assert "COURSE STUDIO" in product_js.text
        assert polish_js.status_code == 200
        assert "运营附加数据" in polish_js.text
        assert studio_js.status_code == 200
        assert "SCRIPT REVIEW" in studio_js.text
        assert "试听选中内容" in studio_js.text
        assert "纠正读音" in studio_js.text
        assert "LAYOUT STUDIO" in studio_js.text
        assert studio_upgrade_js.status_code == 200
        assert "上传自定义背景" in studio_upgrade_js.text
        assert "master_video_asset_id" in studio_upgrade_js.text
        assert page.status_code == 200
        assert "/saas-theme.css" in page.text
        assert "/course-studio.css" in page.text
        assert "/course-studio-upgrade.css" in page.text
        assert "/course-studio.js" in page.text
        assert "/course-studio-upgrade.js" in page.text


def test_speech_preview_queue_is_tenant_scoped_cached_and_unbilled() -> None:
    with TestClient(app) as client:
        token, _, _ = _register(client, "speech-preview")
        headers = _headers(token)

        def upload(name: str, kind: str) -> dict:
            response = client.post(
                "/api/saas/assets",
                headers=headers,
                data={"kind": kind},
                files={"file": (name, BytesIO(b"test-bytes"), "application/octet-stream")},
            )
            assert response.status_code == 201, response.text
            return response.json()

        ppt = upload("preview.pptx", "ppt")
        video = upload("preview.mp4", "video")
        audio = upload("preview.wav", "audio")
        voice_response = client.post(
            "/api/saas/voices",
            headers=headers,
            json={
                "name": "Preview voice",
                "provider": "cosyvoice",
                "reference_asset_id": audio["id"],
                "transcript": "这是一段参考声音",
                "consent_confirmed": True,
            },
        )
        assert voice_response.status_code == 201, voice_response.text
        voice = voice_response.json()
        avatar_response = client.post(
            "/api/saas/avatars",
            headers=headers,
            json={
                "name": "Preview avatar",
                "master_video_asset_id": video["id"],
                "voice_profile_id": voice["id"],
                "consent_confirmed": True,
            },
        )
        assert avatar_response.status_code == 201, avatar_response.text
        narration = "最后我想问下你们一个问题：今天优优认识了谁？"
        course_response = client.post(
            "/api/saas/courses",
            headers=headers,
            json={
                "title": "Speech preview",
                "ppt_asset_id": ppt["id"],
                "avatar_id": avatar_response.json()["id"],
                "voice_profile_id": voice["id"],
                "script": [{"index": 1, "narration": narration}],
                "settings": {"tts": {"pronunciation_replacements": {"优优": "悠悠"}}},
            },
        )
        assert course_response.status_code == 201, course_response.text
        course = course_response.json()
        remaining_before = client.get("/api/saas/billing/subscription", headers=headers).json()["remaining_seconds"]

        payload = {"slide_index": 1, "scope": "selection", "text": "今天优优认识了谁？"}
        first = client.post(
            f"/api/saas/course-tools/courses/{course['id']}/speech-previews",
            headers=headers,
            json=payload,
        )
        assert first.status_code == 202, first.text
        assert first.json()["status"] == "queued"
        repeated = client.post(
            f"/api/saas/course-tools/courses/{course['id']}/speech-previews",
            headers=headers,
            json=payload,
        )
        assert repeated.status_code == 202, repeated.text
        assert repeated.json()["id"] == first.json()["id"]
        remaining_after = client.get("/api/saas/billing/subscription", headers=headers).json()["remaining_seconds"]
        assert remaining_after == remaining_before

        other_token, _, _ = _register(client, "speech-preview-other")
        hidden = client.get(
            f"/api/saas/course-tools/speech-previews/{first.json()['id']}",
            headers=_headers(other_token),
        )
        assert hidden.status_code == 404

        unsaved_page = client.post(
            f"/api/saas/course-tools/courses/{course['id']}/speech-previews",
            headers=headers,
            json={"slide_index": 1, "scope": "page", "text": narration + "改"},
        )
        assert unsaved_page.status_code == 409
